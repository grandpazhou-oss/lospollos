"""Focused legacy boundary checks; subprocess fault injections are labelled separately."""
import asyncio, base64, hashlib, importlib.util, json, os, pathlib, shutil, subprocess, sys, tempfile, unittest
from unittest.mock import patch
ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'backend'))
spec = importlib.util.spec_from_file_location('legacy_server', ROOT / 'backend/server.py')
server = importlib.util.module_from_spec(spec); spec.loader.exec_module(server)
from public_resources import public_file, PUBLIC_FILES
import openpyxl


class Upload:
    def __init__(self, name, content=b'controlled inert bytes'): self.filename, self.content = name, content
    async def read(self, limit=-1): return self.content[:limit] if limit >= 0 else self.content


def value(response):
    return json.loads(response.body) if hasattr(response, 'body') else response


def workbook(path):
    book=openpyxl.Workbook();book.remove(book.active)
    for index in range(7):
        ws=book.create_sheet('Synthetic '+str(index));ws.cell(12,33,'')
    p=book.worksheets[0]
    for col,val in {1:'P1',2:'Synthetic parcel',3:1000,4:900,5:'20261003',8:1,14:2,25:1,26:.5}.items():p.cell(7,col,val)
    c=book.worksheets[1]
    for col,val in {2:2,3:'Synthetic outlet',5:'Fictional address',15:900,16:1730,32:1,33:.5}.items():c.cell(11,col,val)
    co=book.worksheets[3]
    for row,values in [(2,[1,'Synthetic depot','Fictional depot',31,121,'from']),(3,[2,'Synthetic outlet','Fictional address',31.01,121.01,'to'])]:
        for col,val in enumerate(values,1):co.cell(row,col,val)
    ve=book.worksheets[6]
    for col,val in {2:'V1',3:'Synthetic van',5:.6,6:1,7:30,8:6000,14:1,16:900,17:1730}.items():ve.cell(10,col,val)
    book.save(path)


class Boundaries(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(dir=os.environ.get('STCT_SECURITY_TEMP'));self.root=pathlib.Path(self.temp.name);self.old=server.OPTIMIZER_SCRIPT
        self.script=self.root/'delivery_optimizer.py';shutil.copy2(ROOT/'backend/delivery_optimizer.py',self.script);server.OPTIMIZER_SCRIPT=self.script
    def tearDown(self): server.OPTIMIZER_SCRIPT=self.old;self.temp.cleanup()
    def run_upload(self, name, run):
        with patch.object(server.subprocess, 'run', side_effect=run): return asyncio.run(server.upload_dispatch(Upload(name)))
    def test_public_paths_and_aliases(self):
        for name in PUBLIC_FILES:self.assertEqual(public_file('/'+name+'?v=1'),ROOT/name)
        for name in ['/data/uc-study-package.json','/data/routes-data.js','/backend/route_data.js','/backend/route_data.json','/.run/local-trial/web.log','/templates/laiyifen-202605-rawdata.xlsx','/backend/server.py','/vendor/','/../index.html','/%2e%2e/index.html','/data%2fuc-study-package.json','/index.html%00','/index.html%5c','/C:/index.html','/%2564ata/uc-study-package.json']:
            self.assertIsNone(public_file(name),name)
        self.assertEqual(public_file('/'),ROOT/'index.html')
        d=self.root/'public';d.mkdir();(d/'index.html').symlink_to(ROOT/'data/uc-study-package.json');self.assertIsNone(public_file('/index.html',d))
    def test_controlled_subprocess_success_untrusted_filenames_cleanup(self):
        for name in ['../../escape.xlsx','/absolute.xlsx','..\\escape.XLSX',"中文 '空格.xlsx"]:
            seen={}
            def run(args,**kwargs):
                inp,out=pathlib.Path(args[2]),pathlib.Path(args[3]);seen['root']=out
                self.assertEqual(inp.name,'input.xlsx');self.assertEqual(inp.parent,out);self.assertEqual(args[0],sys.executable)
                (out/'route_data.js').write_text('var ROUTE_DATA = {};');(out/'route_data.json').write_text(json.dumps({'routes_by_date':{},'depot':{}}));(out/'delivery_plan.xlsx').write_bytes(b'controlled workbook output')
                return subprocess.CompletedProcess(args,0,'controlled fake subprocess','')
            r=value(self.run_upload(name,run));self.assertTrue(r['success']);self.assertFalse(seen['root'].exists());self.assertFalse((self.root.parent/'escape.xlsx').exists())
            bundle=server.current_legacy_result();self.assertEqual(base64.b64decode(bundle['workbook']),b'controlled workbook output')
            self.assertEqual(value(asyncio.run(server.data_status()))['excelSize'],26)
    def test_failure_timeout_missing_output_keep_current(self):
        current=self.root/'legacy-current.json';current.write_text(json.dumps({'routeData':{'routes_by_date':{}},'workbook':None}));before=current.read_bytes();made=[]
        def fail(args,**kw):made.append(pathlib.Path(args[3]));return subprocess.CompletedProcess(args,1,'','controlled nonzero')
        self.assertFalse(value(self.run_upload('ok.xlsx',fail))['success'])
        def timeout(args,**kw):made.append(pathlib.Path(args[3]));raise subprocess.TimeoutExpired(args,300)
        self.assertFalse(value(self.run_upload('ok.xlsx',timeout))['success'])
        def missing(args,**kw):made.append(pathlib.Path(args[3]));return subprocess.CompletedProcess(args,0,'','')
        self.assertFalse(value(self.run_upload('ok.xlsx',missing))['success']);self.assertEqual(current.read_bytes(),before);self.assertTrue(all(not p.exists() for p in made));self.assertEqual(list(self.root.glob('legacy-result-*.tmp')),[])
    def test_format_and_byte_limit(self):
        for name in [None,'no.csv','no']:
            self.assertFalse(value(asyncio.run(server.upload_dispatch(Upload(name))))['success'])
        r=asyncio.run(server.upload_dispatch(Upload('huge.xlsx',b'x'*(8*1024*1024+1))));self.assertEqual(r.status_code,413)
    def test_native_legacy_optimizer_upload_and_exports(self):
        path=self.root/'合成 小案例.xlsx';workbook(path)
        r=value(asyncio.run(server.upload_dispatch(Upload(path.name,path.read_bytes()))));self.assertTrue(r['success'],r)
        data=value(asyncio.run(server.legacy_route_data()));self.assertEqual(sum(len(v['stops']) for day in data['routes_by_date'].values() for v in day.values()),1)
        export=asyncio.run(server.legacy_export('xlsx'));self.assertEqual(export.status_code,200);self.assertTrue(export.body.startswith(b'PK'))
        self.assertEqual(asyncio.run(server.legacy_export('invalid')).status_code,400)
    def test_cli_literal_argv_and_complete_local_build(self):
        app=self.root/'app';shutil.copytree(ROOT/'backend',app/'backend',ignore=shutil.ignore_patterns('__pycache__','*.json','*.xlsx','*.js'));shutil.copytree(ROOT/'vendor',app/'vendor');shutil.copy2(ROOT/'dispatch.html',app/'dispatch.html')
        path=self.root/"中文 space 'quote\".xlsx";workbook(path)
        result=subprocess.run([sys.executable,str(app/'backend/update.py'),str(path)],capture_output=True,text=True,timeout=120)
        self.assertEqual(result.returncode,0,result.stderr);self.assertTrue((app/'backend/route_data.json').is_file())
        html=(app/'dispatch.html').read_text();self.assertIn('exception-table',html);self.assertIn('vehicle-table',html);self.assertIn('escapeHTML(p.name)',html);self.assertIn('/api/legacy-route-data',html);self.assertNotIn('{{ML_JS}}',html)


if __name__=='__main__':unittest.main(verbosity=2)
